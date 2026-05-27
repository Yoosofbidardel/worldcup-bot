import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import database as db
from config import EXCEL_DIR

HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2E4057")
GOLD_FILL = PatternFill("solid", fgColor="FFD700")
SILVER_FILL = PatternFill("solid", fgColor="C0C0C0")
BRONZE_FILL = PatternFill("solid", fgColor="CD7F32")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center")


def _style_header(ws, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 30)


def generate_excel(chat_id: int) -> str:
    """Generate an Excel report for a group and return the file path."""
    os.makedirs(EXCEL_DIR, exist_ok=True)
    filepath = os.path.join(EXCEL_DIR, f"wc2026_group_{chat_id}.xlsx")
    wb = Workbook()

    # ── Sheet 1: Leaderboard ─────────────────────────────────────
    ws_lb = wb.active
    ws_lb.title = "Leaderboard"
    ws_lb.append(["Rank", "Player", "Username", "Total Points", "Matches Scored"])
    _style_header(ws_lb, 5)

    rows = db.get_leaderboard(chat_id)
    for i, r in enumerate(rows, 1):
        name = r["first_name"] or "Unknown"
        uname = f"@{r['username']}" if r["username"] else ""
        ws_lb.append([i, name, uname, r["total_points"], r["matches_scored"]])
        row_idx = i + 1
        for col in range(1, 6):
            ws_lb.cell(row=row_idx, column=col).border = THIN_BORDER
            ws_lb.cell(row=row_idx, column=col).alignment = CENTER
        if i == 1:
            for col in range(1, 6):
                ws_lb.cell(row=row_idx, column=col).fill = GOLD_FILL
        elif i == 2:
            for col in range(1, 6):
                ws_lb.cell(row=row_idx, column=col).fill = SILVER_FILL
        elif i == 3:
            for col in range(1, 6):
                ws_lb.cell(row=row_idx, column=col).fill = BRONZE_FILL
    _auto_width(ws_lb)

    # ── Sheet 2: All Predictions Detail ──────────────────────────
    ws_pd = wb.create_sheet("Predictions")
    ws_pd.append([
        "Player", "Username", "Match", "Date", "Stage",
        "Prediction", "Actual", "Status", "Points",
    ])
    _style_header(ws_pd, 9)

    preds = db.get_all_predictions_for_export(chat_id)
    for r in preds:
        match_str = f"{r['home_team']} vs {r['away_team']}"
        pred_str = f"{r['pred_home']}-{r['pred_away']}"
        actual_str = f"{r['actual_home']}-{r['actual_away']}" if r["actual_home"] is not None else "-"
        pts = r["points"] if r["points"] is not None else "-"
        ws_pd.append([
            r["first_name"], f"@{r['username']}" if r["username"] else "",
            match_str, r["match_date"][:16], r["stage"],
            pred_str, actual_str, r["status"], pts,
        ])
    _auto_width(ws_pd)

    # ── Sheet 3: Match Results ───────────────────────────────────
    ws_mr = wb.create_sheet("Match Results")
    ws_mr.append(["#", "Home", "Away", "Score", "Date", "Stage", "Group", "Status"])
    _style_header(ws_mr, 8)

    matches = db.get_all_matches()
    for i, m in enumerate(matches, 1):
        score = f"{m['home_score']}-{m['away_score']}" if m["home_score"] is not None else "-"
        ws_mr.append([
            i, m["home_team"], m["away_team"], score,
            m["match_date"][:16], m["stage"], m["group_name"] or "", m["status"],
        ])
    _auto_width(ws_mr)

    # ── Sheet 4: Special Predictions ─────────────────────────────
    ws_sp = wb.create_sheet("Special Predictions")
    ws_sp.append(["Player", "Username", "Type", "Prediction", "Points"])
    _style_header(ws_sp, 5)

    for pred_type in ("top_scorer", "champion"):
        specials = db.get_all_specials(chat_id, pred_type)
        for s in specials:
            ws_sp.append([
                s["first_name"], f"@{s['username']}" if s["username"] else "",
                pred_type.replace("_", " ").title(), s["pred_value"], s["points"],
            ])
    _auto_width(ws_sp)

    wb.save(filepath)
    return filepath
